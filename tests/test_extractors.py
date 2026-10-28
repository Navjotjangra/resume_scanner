"""Tests for scanner/extractors.py per TRD §5.3 and §7.

All fixtures are built in tmp_path (python-docx for .docx incl. a table,
PyMuPDF for .pdf, plain files for txt). No binary files are committed.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from scanner.errors import (
    EmptyTextError,
    ExtractionError,
    FileTooLargeError,
    UnsupportedFormatError,
)
from scanner.extractors import extract_text


def test_extract_txt_utf8(tmp_path: Path) -> None:
    p = tmp_path / "a.txt"
    p.write_text(
        "Python developer with Django experience. Built REST APIs and "
        "worked with PostgreSQL and Docker.",
        encoding="utf-8",
    )
    text = extract_text(p)
    assert "Django" in text


def test_extract_txt_cp1252(tmp_path: Path) -> None:
    p = tmp_path / "b.txt"
    content = (
        "Café résumé — naïve. Experienced software engineer with a degree "
        "in computer science."
    )
    p.write_bytes(content.encode("cp1252"))
    text = extract_text(p)
    assert "Café" in text


def test_extract_txt_utf16_bom(tmp_path: Path) -> None:
    p = tmp_path / "c.txt"
    p.write_bytes(
        "UTF-16 resume content here. Data analyst with SQL and Python skills."
        .encode("utf-16")
    )
    text = extract_text(p)
    assert "resume" in text


def test_extract_docx_body_and_table(tmp_path: Path) -> None:
    from docx import Document

    p = tmp_path / "d.docx"
    doc = Document()
    doc.add_paragraph("Python developer with Django experience.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, SQL"
    table.cell(1, 0).text = "Years"
    table.cell(1, 1).text = "5"
    doc.save(str(p))

    text = extract_text(p)
    assert "Python developer" in text
    assert "Python, SQL" in text


def test_extract_pdf_text(tmp_path: Path) -> None:
    import fitz

    p = tmp_path / "e.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "Machine learning engineer resume. Experienced with TensorFlow, "
        "PyTorch and scikit-learn.",
    )
    doc.save(str(p))
    doc.close()

    text = extract_text(p)
    assert "Machine learning" in text


def test_empty_pdf_raises_empty_text(tmp_path: Path) -> None:
    import fitz

    p = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()  # blank page, no extractable text
    doc.save(str(p))
    doc.close()

    with pytest.raises(EmptyTextError):
        extract_text(p)


def test_unsupported_extension(tmp_path: Path) -> None:
    p = tmp_path / "x.xyz"
    p.write_bytes(b"garbage")
    with pytest.raises(UnsupportedFormatError):
        extract_text(p)


def test_missing_file(tmp_path: Path) -> None:
    p = tmp_path / "does_not_exist.txt"
    with pytest.raises(ExtractionError):
        extract_text(p)


def test_oversize_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("scanner.extractors.MAX_FILE_MB", 0)
    p = tmp_path / "big.txt"
    p.write_text("some content here")
    with pytest.raises(FileTooLargeError):
        extract_text(p)


def test_corrupt_pdf_raises_extraction_error(tmp_path: Path) -> None:
    p = tmp_path / "corrupt.pdf"
    p.write_bytes(b"this is definitely not a valid pdf file")
    with pytest.raises(ExtractionError):
        extract_text(p)


def test_doc_without_soffice_raises_unsupported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    p = tmp_path / "legacy.doc"
    p.write_bytes(b"fake legacy doc content")
    monkeypatch.setattr("scanner.extractors.shutil.which", lambda *_: None)
    monkeypatch.setattr("scanner.extractors.SOFFICE_CANDIDATES", ())
    with pytest.raises(UnsupportedFormatError) as excinfo:
        extract_text(p)
    assert "LibreOffice" in str(excinfo.value)

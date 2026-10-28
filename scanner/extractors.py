"""Text extraction for PDF, DOCX, DOC and TXT files.

Dispatches by extension and raises scanner.errors exceptions so the
pipeline can collect per-file failures without crashing the scan.
"""

from __future__ import annotations

import codecs
import shutil
import subprocess
import tempfile
from pathlib import Path

from scanner.config import (
    MAX_FILE_MB,
    MIN_TEXT_CHARS,
    SOFFICE_CANDIDATES,
    SOFFICE_TIMEOUT_S,
    SUPPORTED_EXTENSIONS,
)
from scanner.errors import (
    EmptyTextError,
    ExtractionError,
    FileTooLargeError,
    UnsupportedFormatError,
)

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover - fallback for older installs
    import fitz  # type: ignore[no-redef]


def extract_text(path: Path) -> str:
    """Extract readable text from a supported file.

    Steps, in order:
    1. Path must exist and be a file, else ExtractionError.
    2. Extension (lower-case) must be in SUPPORTED_EXTENSIONS, else
       UnsupportedFormatError.
    3. Size above MAX_FILE_MB -> FileTooLargeError.
    4. Dispatch by extension; wrap library errors as ExtractionError.
    5. Strip; fewer than MIN_TEXT_CHARS chars -> EmptyTextError.
    """
    path = Path(path)
    if not path.exists() or not path.is_file():
        raise ExtractionError(f"File not found: {path.name}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"Unsupported file type: {ext or 'no extension'}"
        )

    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > MAX_FILE_MB:
        raise FileTooLargeError(
            f"File larger than {MAX_FILE_MB} MB: {path.name}"
        )

    if ext == ".pdf":
        text = _extract_pdf(path)
    elif ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".doc":
        text = _extract_doc(path)
    else:  # .txt
        text = _extract_txt(path)

    text = text.strip()
    if len(text) < MIN_TEXT_CHARS:
        raise EmptyTextError("No readable text (scanned or image-only file?)")

    return text


def _extract_pdf(path: Path) -> str:
    """Extract text from a PDF via PyMuPDF (pymupdf, fitz fallback)."""
    try:
        with fitz.open(path) as doc:
            if doc.needs_pass:
                raise ExtractionError("Password-protected PDF")
            pages = [page.get_text("text") for page in doc]
    except ExtractionError:
        raise
    except Exception as e:  # noqa: BLE001 - wrap any library error
        raise ExtractionError(f"Could not read file: {e}") from e
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    """Extract text from a DOCX: body paragraphs plus all table cells."""
    try:
        from docx import Document

        doc = Document(path)
        parts: list[str] = []

        for paragraph in doc.paragraphs:
            if paragraph.text:
                parts.append(paragraph.text)

        def _collect_tables(table) -> None:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
                    for nested in cell.tables:
                        _collect_tables(nested)

        for table in doc.tables:
            _collect_tables(table)
    except ExtractionError:
        raise
    except Exception as e:  # noqa: BLE001 - wrap any library error
        raise ExtractionError(f"Could not read file: {e}") from e
    return "\n".join(parts)


def _extract_txt(path: Path) -> str:
    """Extract text from a TXT with the TRD encoding fallback order."""
    try:
        raw = path.read_bytes()
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not read file: {e}") from e

    # UTF-16 BOM (little- or big-endian)
    if raw[:2] in (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE):
        return raw.decode("utf-16")

    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    return raw.decode("latin-1", errors="replace")


def _find_soffice() -> str | None:
    """Locate the LibreOffice binary via PATH or known install paths."""
    if soffice := shutil.which("soffice"):
        return soffice
    for candidate in SOFFICE_CANDIDATES:
        if Path(candidate).exists():
            return candidate
    return None


def _extract_doc(path: Path) -> str:
    """Best-effort .doc extraction via LibreOffice conversion to .docx."""
    soffice = _find_soffice()
    if not soffice:
        raise UnsupportedFormatError(
            ".doc needs LibreOffice installed, or save the file as .docx"
        )

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    tmpdir,
                    str(path),
                ],
                timeout=SOFFICE_TIMEOUT_S,
                capture_output=True,
                check=True,
            )
            converted = Path(tmpdir) / (path.stem + ".docx")
            if not converted.exists():
                raise FileNotFoundError("LibreOffice did not produce output")
    except Exception as e:  # noqa: BLE001 - conversion failure
        raise UnsupportedFormatError(
            ".doc needs LibreOffice installed, or save the file as .docx"
        ) from e

    return _extract_docx(converted)
